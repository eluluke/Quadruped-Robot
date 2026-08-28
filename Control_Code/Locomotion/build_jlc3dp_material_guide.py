from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
BUILD_DIR = ROOT / "jlc3dp_material_guide_build"
OUTPUT = ROOT / "JLCPCB_3D_Printing_Material_Guide_for_Quadruped_Robot.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "183247"
ORANGE = "E97824"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F3F7FA"
PALE_ORANGE = "FFF1E6"
PALE_GREEN = "E8F3EC"
PALE_RED = "FBE9E7"
GRAY = "5C6770"
LIGHT_GRAY = "F2F4F6"
WHITE = "FFFFFF"
BLACK = "1F252B"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color="D5DCE2", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_widths(table, widths_inches: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_inches):
            row.cells[idx].width = Inches(width)


def style_run(run, size=None, bold=None, color=None, italic=None, font="Calibri") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    style_run(run, 8, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_table(doc, headers, rows, widths, font_size=8.2, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(value))
        style_run(run, font_size, bold=True, color=WHITE)
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if ridx % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        for cidx, value in enumerate(row_data):
            cell = row.cells[cidx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            style_run(run, font_size, bold=(first_col_bold and cidx == 0), color=BLACK)
    set_table_widths(table, widths)
    return table


def add_callout(doc, title: str, text: str, fill=PALE_BLUE, title_color=DARK_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.45)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 130, 170, 130, 170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    style_run(run, 10.5, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    run = p2.add_run(text)
    style_run(run, 9.5, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text: str, level=0, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.28 + 0.22 * level)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        style_run(run, 10, bold=True, color=DARK_BLUE)
        run = p.add_run(text[len(bold_prefix):])
        style_run(run, 10, color=BLACK)
    else:
        run = p.add_run(text)
        style_run(run, 10, color=BLACK)


def add_label_paragraph(doc, label: str, text: str, fill=None) -> None:
    if fill:
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        set_cell_shading(cell, fill)
        set_cell_margins(cell, 90, 130, 90, 130)
        p = cell.paragraphs[0]
    else:
        p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(label)
    style_run(run, 10, bold=True, color=DARK_BLUE)
    run = p.add_run(text)
    style_run(run, 10, color=BLACK)


def add_scorecard_table(doc: Document) -> None:
    materials = [
        "MJF PA12-HP",
        "MJF PA11-HP",
        "SLS PA12 family",
        "SLS 3401GB",
        "FDM PA12-CF",
        "FDM ASA / ABS",
        "FDM PEEK",
        "SLA engineering resin",
        "TPU / PEBA",
        "SLM TC4 titanium",
        "SLM 316L",
        "BJ 316L",
    ]
    criteria = ["Specific\nstrength", "Stiffness", "Impact /\nfatigue", "Heat", "Process\nconsistency", "Cost\naccess"]
    scores = [
            [4, 3, 4, 4, 5, 4],
            [4, 3, 5, 4, 5, 3],
            [4, 3, 4, 4, 4, 5],
            [3, 4, 2, 4, 4, 4],
            [4, 5, 2, 4, 2, 4],
            [3, 3, 3, 3, 2, 5],
            [4, 4, 2, 5, 2, 1],
            [2, 3, 1, 2, 5, 5],
            [1, 1, 5, 2, 2, 4],
            [5, 5, 4, 5, 4, 1],
            [3, 5, 5, 5, 4, 2],
            [3, 5, 4, 5, 3, 3],
        ]
    table = doc.add_table(rows=1, cols=7)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table, color="FFFFFF", size="8")
    headers = ["Material family"] + criteria
    widths = [1.55, 0.82, 0.73, 0.88, 0.65, 0.95, 0.82]
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell, 70, 50, 70, 50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(value)
        style_run(run, 7.4, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    palette = {1: "F3B6AE", 2: "F8D7B1", 3: "FFF0B8", 4: "CFE7C9", 5: "93C89A"}
    for name, row_scores in zip(materials, scores):
        row = table.add_row()
        prevent_row_split(row)
        name_cell = row.cells[0]
        set_cell_shading(name_cell, PALE_BLUE)
        set_cell_margins(name_cell, 58, 70, 58, 70)
        p = name_cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(name)
        style_run(run, 7.8, bold=True, color=NAVY)
        for idx, score in enumerate(row_scores, 1):
            cell = row.cells[idx]
            set_cell_shading(cell, palette[score])
            set_cell_margins(cell, 58, 30, 58, 30)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(score))
            style_run(run, 8.4, bold=True, color=NAVY)
    set_table_widths(table, widths)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = True
    section.different_first_page_header_footer = False
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 10),
        ("Subtitle", 13, GRAY, 0, 8),
        ("Heading 1", 17, BLUE, 14, 8),
        ("Heading 2", 13, DARK_BLUE, 11, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Small Note" not in styles:
        note = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = styles["Small Note"]
    note.font.name = "Calibri"
    note._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    note.font.size = Pt(8.5)
    note.font.color.rgb = RGBColor.from_string(GRAY)
    note.paragraph_format.space_after = Pt(4)
    note.paragraph_format.line_spacing = 1.05

    for header in (section.header, section.even_page_header):
        table = header.add_table(rows=1, cols=2, width=Inches(6.45))
        table.columns[0].width = Inches(3.55)
        table.columns[1].width = Inches(2.90)
        table.autofit = False
        p = table.cell(0, 0).paragraphs[0]
        run = p.add_run("JLCPCB / JLC3DP MATERIAL GUIDE")
        style_run(run, 8, bold=True, color=BLUE)
        p = table.cell(0, 1).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run("QUADRUPED ROBOT DESIGN")
        style_run(run, 8, bold=True, color=GRAY)

    for footer in (section.footer, section.even_page_footer):
        table = footer.add_table(rows=1, cols=2, width=Inches(6.45))
        table.columns[0].width = Inches(4.55)
        table.columns[1].width = Inches(1.90)
        table.autofit = False
        p = table.cell(0, 0).paragraphs[0]
        run = p.add_run("Prepared 26 August 2026 | Verify live quote and current TDS before release")
        style_run(run, 8, color=GRAY)
        add_page_number(table.cell(0, 1).paragraphs[0])


def add_source(doc, number: int, title: str, url: str, note: str) -> None:
    p = doc.add_paragraph(style="Small Note")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    run = p.add_run(f"[{number}] ")
    style_run(run, 8.5, bold=True, color=DARK_BLUE)
    add_hyperlink(p, title, url)
    run = p.add_run(f". {note}")
    style_run(run, 8.5, color=GRAY)


def build_document() -> None:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_document(doc)
    core = doc.core_properties
    core.title = "JLCPCB / JLC3DP 3D Printing Material Selection for a Quadruped Robot"
    core.subject = "Engineering comparison of JLC3DP materials and component-level recommendations"
    core.author = "Quadruped Robot Project"
    core.keywords = "JLCPCB, JLC3DP, quadruped robot, PA12, PA11, PA12-CF, SLS, MJF, FDM, SLA, SLM"

    # Cover
    stripe = doc.add_table(rows=1, cols=1)
    stripe.cell(0, 0).height = Inches(0.10)
    set_cell_shading(stripe.cell(0, 0), ORANGE)
    set_cell_margins(stripe.cell(0, 0), 0, 0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("ENGINEERING MATERIAL SELECTION GUIDE")
    style_run(run, 9.5, bold=True, color=ORANGE)

    p = doc.add_paragraph(style="Title")
    p.add_run("JLCPCB / JLC3DP\n3D Printing Materials")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Strength, cost, manufacturability, and component recommendations for a quadruped robot")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    add_callout(
        doc,
        "Decision in one sentence",
        "Use MJF PA12-HP or SLS PA12-family nylon for most load-bearing body and leg structures; reserve PA12-CF for stiffness-dominated parts that can tolerate FDM anisotropy, PA11 for impact-dominated parts, TPU/PEBA for feet, and printed metal only for compact interfaces where polymer is demonstrably insufficient.",
        fill=PALE_ORANGE,
        title_color=ORANGE,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Prepared for")
    style_run(run, 9, bold=True, color=GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("Quadruped Robot Project")
    style_run(run, 14, bold=True, color=NAVY)
    p = doc.add_paragraph(style="Small Note")
    p.add_run("Research date: 26 August 2026 | Catalog availability and prices can change")

    p = doc.add_paragraph(style="Small Note")
    p.paragraph_format.space_before = Pt(16)
    p.add_run(
        "Scope: Publicly listed JLC3DP/JLCPCB additive-manufacturing materials and service specifications. "
        "The document distinguishes published values from engineering screening judgments. It is not a substitute for coupon testing, FEA, fatigue testing, or a supplier-certified material lot report."
    )

    doc.add_page_break()

    # Executive recommendation
    doc.add_heading("1. Executive Recommendation", level=1)
    add_callout(
        doc,
        "Recommended material stack",
        "A mixed-material robot is the technically strongest and usually the cheapest solution. Do not force one material to serve as structure, bearing interface, heat shield, cable protector, and foot tread.",
    )

    recommendation_rows = [
        ("Main body / chassis shells", "MJF PA12-HP", "Best overall balance of toughness, fatigue resistance, dimensional repeatability, complex geometry, and low mass.", "SLS 3301PA or 3201PA-F when quote cost or size is better."),
        ("Thigh and shank link shells", "MJF PA12-HP or SLS 3301PA", "Tough powder-bed nylon tolerates impacts and has less directional weakness than FDM.", "PA12-CF only when stiffness is the dominant problem and orientation is controlled."),
        ("Impact-prone clips / compliant brackets", "MJF PA11-HP", "JLC3DP publishes 50% elongation at break and equal XY/Z tensile strength.", "PA12-HP for lower cost or better dimensional familiarity."),
        ("Actuator brackets / bearing carriers", "TC4 SLM for weight-critical; 316L SLM for cost-tolerant mass", "High strength and better creep resistance at concentrated loads.", "Powder nylon with metal sleeves only after load testing."),
        ("Gearbox housings / precision bores", "Metal interface or PA12-CF shell plus machined metal inserts", "Printed polymer alone is vulnerable to creep, bore growth, and bearing-seat damage.", "MJF PA12 with large wall sections for low-load prototypes."),
        ("Foot pads / bumpers / strain relief", "TPU 95A", "Abrasion resistance, replaceability, and compliance.", "PEBA 85A when elastic energy return matters more than damping."),
        ("Electronics covers / outdoor trim", "ASA", "UV resistance and near-ABS heat performance.", "ABS indoors; SLA for visual and fit checks only."),
        ("Cycloidal discs / loaded gear teeth", "Machined metal preferred", "High-cycle contact stress and wear are poor matches for most printed polymers.", "PA12-CF or PEEK only for controlled low-load experiments."),
    ]
    add_table(
        doc,
        ["Component", "Primary choice", "Why", "Alternate / boundary"],
        recommendation_rows,
        [1.30, 1.25, 2.25, 1.65],
        font_size=8.3,
        first_col_bold=True,
    )

    doc.add_heading("What not to use for primary structure", level=2)
    add_bullet(doc, "Standard SLA resins for cyclic thigh, shank, chassis, or gearbox loads. Their clean surface and high tensile numbers do not remove UV aging, brittle fracture, micro-cracking, and poor impact tolerance.")
    add_bullet(doc, "PLA-P for motor-adjacent or load-bearing robot parts. JLC3DP lists a 65 °C HDT and explicitly excludes prolonged heat and load-bearing use.")
    add_bullet(doc, "Full-color PAC-HP or WJP resin for structural legs. Their purpose is appearance and detailed color, not high-cycle mechanics.")
    add_bullet(doc, "As-printed holes or printed threads at motor mounts, bearings, pivots, or battery restraints. Use metal sleeves, through-bolts, inserts, or post-machining.")

    doc.add_heading("2. How to Read the Data", level=1)
    property_rows = [
        ("Tensile strength", "Peak stress before tensile failure", "Useful, but not a direct predictor of joint survival under bending, impact, fatigue, or layer separation."),
        ("Elastic / tensile modulus", "Resistance to elastic deflection", "Higher modulus makes links hold geometry better, but may reduce toughness and increase shock transmission."),
        ("Elongation at break", "Ductility before fracture", "A practical warning signal: PA11 at 50% is damage-tolerant; PA12-CF at 2.9% is much less forgiving."),
        ("Impact strength", "Resistance to sudden energy input", "Important for foot strikes, falls, backlash impacts, and joint hard stops."),
        ("HDT", "Temperature at which a loaded specimen deflects", "Always read the test load. PA11 is 185 °C at 0.45 MPa but only 54 °C at 1.82 MPa."),
        ("Process anisotropy", "Direction-dependent properties", "FDM Z-strength can be far below XY strength; powder-bed processes are generally more uniform."),
        ("Creep / fatigue", "Time- and cycle-dependent deformation", "Usually more important than one-time tensile strength for a walking robot; supplier tables rarely capture your exact duty cycle."),
    ]
    add_table(doc, ["Property", "Meaning", "Why it matters on the robot"], property_rows, [1.35, 1.70, 3.40], font_size=8.6, first_col_bold=True)
    add_label_paragraph(
        doc,
        "Comparison rule: ",
        "Do not rank materials using one number from different test standards, orientations, moisture conditions, wall thicknesses, or infill settings. Use the published values for screening, then test production-intent coupons and a representative link or bracket.",
        fill=PALE_ORANGE,
    )

    # Scorecard
    doc.add_heading("3. Engineering Screening Scorecard", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "The heatmap converts the published data and process behavior into a quick 1-to-5 screening score. "
        "It is a design aid, not a supplier guarantee. Cost access includes process cost, likely finishing, and the risk of failed iterations."
    )
    add_scorecard_table(doc)
    p = doc.add_paragraph(style="Small Note")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Table 1. Relative quadruped-component screening score. Compare within a row and validate against the real load case.")

    doc.add_heading("Interpretation", level=2)
    add_bullet(doc, "MJF PA12-HP is the safest general-purpose structural polymer; SLS PA12-family grades are close and often cost-effective.")
    add_bullet(doc, "PA11-HP is the toughness leader; 3401GB and PA12-CF are stiffness-first choices with lower ductility.")
    add_bullet(doc, "PEEK solves heat and chemical problems, not automatically impact or layer-bonding problems. JLC3DP publishes only 25.6 MPa tensile strength in Z versus 73 MPa in XY.")
    add_bullet(doc, "Titanium offers the best printed-metal strength-to-weight ratio; 316L is heavier but more ductile and corrosion resistant.")

    # Process comparison
    doc.add_heading("4. Process-Level Comparison", level=1)
    process_rows = [
        ("SLA resin", "Very smooth; ±0.2 mm within 100 mm", "Supports required; easy fine detail", "UV/heat aging and brittle cyclic behavior", "$–$$", "Fit checks, molds, covers, visual prototypes"),
        ("MJF nylon", "Consistent, fine grain; ±0.3 mm", "No support structures; good batch repeatability", "Powder texture; moisture and creep still matter", "$$", "Primary body and leg structures"),
        ("SLS nylon", "Grainier; ±0.3 mm", "No supports; complex hollow and ribbed geometry", "Surface finish and dimensions less crisp than SLA", "$–$$", "Structural links, housings, prototypes"),
        ("FDM plastic", "Visible layers; ±0.3 mm", "Large parts, familiar design rules", "Strong anisotropy, warping, infill dependence", "$–$$$", "Covers, prototypes, stiffness-first PA12-CF parts"),
        ("SLM metal", "Rough; ±0.3 mm", "Near-dense metal, complex topology", "High cost, supports, stress relief and machining", "$$$$", "Compact high-load brackets and interfaces"),
        ("Binder-jet metal", "More even than SLM; tolerance degrades >50 mm", "No melt support strategy; economical batches", "Sinter shrinkage, lower yield strength, 100 mm build limit", "$$$", "Small metal fittings, not precision joint geometry"),
        ("WJP full color", "Excellent color/detail; ±0.2 mm", "Appearance-first workflow", "Low heat/UV suitability and high cost", "$$$", "Team models, presentation and anatomy demonstrators"),
    ]
    add_table(
        doc,
        ["Process", "Accuracy / finish", "Manufacturability", "Main engineering limit", "Cost", "Robot fit"],
        process_rows,
        [0.78, 1.08, 1.32, 1.42, 0.48, 1.35],
        font_size=7.8,
        first_col_bold=True,
    )
    p = doc.add_paragraph(style="Small Note")
    p.add_run("JLC3DP publishes general process tolerances of ±0.2 mm for SLA/WJP and ±0.3 mm for MJF, SLS, FDM, and SLM within 100 mm. Holes usually print undersize and require machining allowance.")

    # Nylon catalog
    doc.add_heading("5. Powder-Bed Nylon Materials", level=1)
    p = doc.add_paragraph()
    p.add_run("These are the most relevant materials for the robot's body and legs because they combine low density, toughness, complex geometry, and relatively uniform properties without FDM support structures.")
    nylon_rows = [
        ("PA12-HP", "MJF", "$1 / $$", "48 MPa | 1.8 GPa | 20%", "175 °C @0.45 MPa", "Best all-round structural choice; strong, durable, chemically resistant."),
        ("PA11-HP", "MJF", "$1 / $$", "52 MPa | 1.8 GPa | 50%", "185 °C @0.45; 54 °C @1.82 MPa", "Highest ductility and impact tolerance; excellent for clips and shock-prone links."),
        ("PAC-HP", "MJF full color", "$17 / $$$", "No public strength table", "80 °C", "Color/detail material for models and craft prototypes, not a primary structure."),
        ("3301PA", "SLS", "$1 / $–$$", "48 MPa | 1.7 GPa | 30%", "147 °C @0.45; 55 °C @1.8 MPa", "Tough white structural nylon; strong candidate for body and links."),
        ("3201PA-F", "SLS", "$1 / $–$$", "44 MPa | 1.5 GPa | 35%", "147 °C @0.45 MPa", "Very tough, accurate, economical structural prototype and end-use option."),
        ("1172 Pro", "SLS", "$1 / $–$$", "46 MPa | 1.8 GPa | 8–15%", "179 °C @0.45 MPa", "Good general mechanics and tall 590 mm build; lower ductility than PA12-HP/PA11."),
        ("3401GB", "SLS glass-bead", "$1 / $$", "42 MPa | 2.5 GPa | 8%", "150 °C @0.45; 60 °C @1.8 MPa", "Stiffer and dimensionally stable; less impact tolerant. Use for housings/supports, not thin shock-loaded links."),
        ("PA12S-HP", "MJF; support table", "Quote / unknown", "Dedicated public TDS not located", "Verify current TDS", "Appears in JLC3DP design tables; do not substitute for PA12-HP without the live material sheet."),
    ]
    add_table(
        doc,
        ["Material", "Process", "Entry / relative cost", "Tensile | modulus | elong.", "Published HDT", "Engineering judgment"],
        nylon_rows,
        [0.85, 0.72, 0.86, 1.30, 1.05, 1.80],
        font_size=7.7,
        first_col_bold=True,
    )
    p = doc.add_paragraph(style="Small Note")
    p.add_run("Cost notation: the first value is JLC3DP's advertised starting price, not the final part quote. Relative tiers include process and finishing burden. JLC3DP's capability page contains a '3021PA-F' label while its material page and design guide use '3201PA-F'; verify the live quote name before ordering.")

    doc.add_heading("Nylon recommendation hierarchy", level=2)
    add_bullet(doc, "Default: PA12-HP MJF for the main chassis and most leg shells.")
    add_bullet(doc, "Toughness first: PA11-HP where falls, clips, flexing, or repeated shock dominate.")
    add_bullet(doc, "Cost/large geometry: 3301PA or 3201PA-F SLS.")
    add_bullet(doc, "Stiffness first: 3401GB for stable housings and supports, with generous radii and no sharp shock path.")

    # FDM catalog
    doc.add_heading("6. FDM Plastics", level=1)
    fdm_rows = [
        ("PLA-P", "$1 / $", "36–40 MPa | 2.55–2.68 GPa | 2.5–6%", "65 °C", "Cheap prototypes and jigs only; poor heat and long-term load behavior."),
        ("ABS", "$3 / $", "28.8 MPa | 1.85 GPa | 3.8%", "98 °C", "Indoor covers and thick simple functional parts; layer strength and warping limit structural use."),
        ("ASA", "$3 / $$", "43.8 MPa* | 2.38 GPa | 6.7%", "97.8 °C", "Best FDM outdoor cover material because of UV resistance. *Supplier page appears to mislabel the tensile unit; verify TDS."),
        ("PA12-CF", "$3 / $$", "69.3 MPa | 3.75 GPa | 2.9%", "105 °C", "Very stiff and strong in favorable orientation; low elongation and FDM anisotropy demand coupons and conservative joints."),
        ("TPU 95A", "$3 / $$", "No public tensile table on service page", "Not listed", "Replaceable foot pads, bumpers, cable strain relief; not a rigid structure."),
        ("PEBA 85A", "$3 / $$$", "21 MPa | 0.063 GPa | 600%", "Low-temp capability to -70 °C", "High elastic return for feet/lattices; less damping and harder process than TPU."),
        ("PEEK", "$3 entry / $$$$ actual", "XY 73, Z 25.6 MPa | ~3.0 GPa | 3.7/1.3%", "144 °C @1.8 MPa", "High-temperature/chemical specialty option; pronounced Z weakness and expensive process development."),
        ("ABS-ESD", "$3 / $$", "XY 38, Z 16 MPa | 2.0–2.5 GPa | 1.5–3.9%", "Not listed", "ESD-safe electronics fixtures and enclosures; not worth the strength penalty for ordinary robot structure."),
    ]
    add_table(
        doc,
        ["Material", "Entry / relative cost", "Tensile | modulus | elong.", "Thermal marker", "Recommended role / limit"],
        fdm_rows,
        [0.92, 1.00, 1.62, 1.15, 2.00],
        font_size=8.0,
        first_col_bold=True,
    )
    add_callout(
        doc,
        "Why PA12-CF is not automatically the winner",
        "Its 3.75 GPa modulus is roughly twice PA12-HP's published 1.8 GPa, so it can reduce link flex. But its 2.9% elongation is far below PA12-HP's 20% and PA11-HP's 50%, and FDM layer interfaces remain orientation-dependent. Use it for a stiff shell with aligned load paths, not for thin snap features, hard-stop impacts, or an untested gearbox disc.",
        fill=PALE_ORANGE,
        title_color=ORANGE,
    )

    # Metals
    doc.add_heading("7. Printed Metals", level=1)
    metal_rows = [
        ("TC4 titanium", "SLM", "$8 / $$$$", "1160 MPa tensile; 1000 MPa yield; 8% elong.", "Best printed strength-to-weight option; rough and difficult to machine.", "Weight-critical actuator brackets, pin carriers, compact high-load interfaces."),
        ("316L stainless", "SLM", "$8 / $$$", "600 MPa tensile; 400 MPa yield; 30% elong.; E=180 GPa", "Strong, ductile, corrosion resistant, but heavy; bores need post-machining.", "Small bearing carriers, brackets, threaded bosses, safety-critical mounts."),
        ("BJ-316L", "Binder jet", "$5 / $$$", "561 MPa tensile; 219 MPa yield; 50% elong.", "Cheaper and smoother; shrink/warp risk and ±1.3% tolerance above 50 mm.", "Small non-precision fittings and low-volume metal parts."),
    ]
    add_table(
        doc,
        ["Material", "Process", "Entry / relative cost", "Published mechanics", "Trade-off", "Robot use"],
        metal_rows,
        [0.90, 0.68, 0.88, 1.35, 1.55, 1.48],
        font_size=8.0,
        first_col_bold=True,
    )
    add_label_paragraph(
        doc,
        "Weight warning: ",
        "A steel bracket may solve a local stress problem while worsening whole-body inertia and leg torque demand. Prefer topology-optimized titanium only where the load path justifies the price; otherwise compare a machined aluminum part before choosing printed stainless steel.",
        fill=PALE_RED,
    )

    # Resins
    doc.add_heading("8. SLA and Full-Color Resins", level=1)
    p = doc.add_paragraph()
    p.add_run("JLC3DP offers many resin names, but the structural differences are smaller than the gap between resin and powder-bed nylon. Treat most as precision prototype materials unless a dedicated fatigue and environmental test says otherwise.")
    resin_rows_1 = [
        ("Grey Resin", "$0.30", "HDT 51–59 °C; tensile 41–61 MPa; elong. 7–12%", "Low-cost gray fit/appearance and light-duty prototype."),
        ("Black Resin", "$1", "HDT 58–69 °C; tensile 41–58 MPa; elong. 7–11%", "Smooth black prototype; no UV/outdoor use."),
        ("JLC Black", "$1", "HDT 54–65 °C; supplier emphasizes hardness and finish", "Batch appearance/functional prototypes, not cyclic legs."),
        ("Imagine Black", "$2", "HDT 46 °C; elong. 6–9%", "High-precision/high-humidity appearance parts; low heat margin."),
        ("9600", "$0.30", "HDT 59 °C; matte white", "Economical white enclosures, pipes, buckles, and fit checks."),
        ("9000HE", "$0.30", "HDT 60 °C; tensile 41–61 MPa; elong. 7–12%", "General high-strength/tough resin prototype; still UV/heat limited."),
        ("8228", "$0.30", "HDT 56.4 °C; tensile 62.4 MPa; elong. 6.8%", "Higher published tensile strength and stable batch prototypes."),
    ]
    add_table(doc, ["Material", "Entry price", "Published differentiator", "Best use / warning"], resin_rows_1, [1.05, 0.72, 2.45, 2.05], font_size=8.2, first_col_bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    resin_rows_2 = [
        ("8001", "$0.30", "Transparent/translucent; HDT 53 °C", "Optical/flow/assembly visualization; default sanding and clear oil finish."),
        ("LEDO 6060", "$1 on JLCPCB listing", "HDT about 56 °C; dimensional stability and water-resistant prototype use", "General housings and waterproof concept verification; not primary structure."),
        ("CBY", "$0.30", "HDT 72–80 °C; tensile 25–30 MPa; elong. 12–20%", "More compliant/high-temperature resin for molds and functional prototypes."),
        ("JLC Temp", "$0.30", "HDT 101 °C; supplier warns it becomes brittle after cure", "Static hot-air/tooling/wind-tunnel parts, not impacts or legs."),
        ("X Resin", "$0.30", "Randomly assigned resin/color", "Cheap non-critical mockups only; impossible to engineer around an unknown grade."),
        ("Full-Color Resin", "$5", "WJP; HDT 60–65 °C; tensile 45–55 MPa; elong. 12–25%", "Team communication models and presentation parts, not final robot structure."),
        ("Alternate names", "Quote-dependent", "JLCPCB pages also show 8000, 9000R, and 8111X", "Treat as SLA prototype grades and verify the exact live TDS before substitution."),
    ]
    add_table(doc, ["Material", "Entry price", "Published differentiator", "Best use / warning"], resin_rows_2, [1.05, 0.72, 2.45, 2.05], font_size=8.2, first_col_bold=True)

    # Design rules
    doc.add_heading("9. Design Rules for Robot Components", level=1)
    design_rows = [
        ("Wall thickness", "Supplier minimum is a printability limit, not a structural recommendation.", "Start structural nylon shells around 3–4 mm with ribs, then validate by FEA and test; thicken locally at joints."),
        ("Fasteners", "Printed threads and point-loaded bosses crack or creep.", "Use through-bolts, metal sleeves, washers, captured nuts, or M3–M5 inserts. Keep generous edge distance."),
        ("Bearings / shafts", "As-printed bores shrink and wear.", "Add machining allowance and use metal bearing cups or sleeves. Do not press a bearing directly into a thin resin wall."),
        ("Load paths", "Sharp corners and sudden section changes concentrate stress.", "Use large fillets, closed sections, ribs, and gradual transitions. Carry actuator load into multiple walls."),
        ("Powder removal", "SLS/MJF powder remains trapped in closed cavities.", "Add escape holes and inspection access; avoid inaccessible internal channels unless required."),
        ("FDM orientation", "Z interfaces are the weakest direction.", "Orient layer planes to carry bending and bolt loads; avoid peeling a flange across layers."),
        ("Thermal exposure", "Motor and controller heat can reduce stiffness long before melting.", "Measure local temperature under worst duty cycle and compare to HDT at the relevant stress, not only the 0.45 MPa number."),
        ("Replaceability", "Feet, hard-stop pads, and cable guards are wear parts.", "Make TPU/PEBA components replaceable without disassembling the whole leg."),
    ]
    add_table(doc, ["Design topic", "Failure mechanism", "Recommended practice"], design_rows, [1.12, 2.15, 3.18], font_size=8.5, first_col_bold=True)

    doc.add_heading("10. Validation Plan Before Full-Robot Use", level=1)
    validation_steps = [
        ("1", "Freeze the load case", "Record robot mass, peak actuator torque, joint geometry, hard-stop condition, fall height, motor temperature, and target life cycles."),
        ("2", "Print coupons in production orientation", "Test at least tensile/bend coupons plus a bolt-boss or bearing-seat coupon. Condition nylon in the expected humidity."),
        ("3", "Test one representative component", "Proof-load a thigh/shank or bracket gradually while measuring deflection. Inspect for white stress marks, cracks, bore growth, and insert movement."),
        ("4", "Cycle under realistic motion", "Run the component through repeated direction changes and impact-like loads. Static success does not establish fatigue life."),
        ("5", "Instrument and inspect", "Mark fasteners, measure temperatures, log joint error/current, and remeasure bores and critical dimensions after cycling."),
        ("6", "Release with a documented margin", "Keep a conservative safety factor for uncertainty and impacts; replace parts on measured wear or cycle count, not only after visible failure."),
    ]
    add_table(doc, ["Step", "Action", "Acceptance intent"], validation_steps, [0.46, 1.55, 4.44], font_size=8.6, first_col_bold=True)
    add_callout(
        doc,
        "Recommended first procurement experiment",
        "Order the same small structural bracket or scaled link section in PA12-HP, PA11-HP, 3301PA, 3401GB, and PA12-CF. Use identical holes and load fixtures. Compare mass, deflection, permanent set, impact damage, thread/insert retention, and final delivered price. That experiment will be more useful than selecting from tensile strength alone.",
        fill=PALE_GREEN,
        title_color="2E6B47",
    )

    # Final decision
    doc.add_heading("11. Final Component Selection", level=1)
    final_rows = [
        ("Body frame and battery/electronics tray", "MJF PA12-HP", "Ribbed shell, metal standoffs, through-bolts; keep heavy battery attachment in a distributed load path."),
        ("Thigh and shank", "MJF PA12-HP or SLS 3301PA", "Closed-section shell with large root fillets and metal pivot sleeves. Choose PA11 if falls and impacts dominate."),
        ("High-stiffness link revision", "FDM PA12-CF", "Only after orientation-specific coupon and full link proof test; avoid thin bosses and snap features."),
        ("Hip/actuator brackets", "TC4 SLM for minimum mass; 316L SLM for lower material cost", "Post-machine bearing and alignment surfaces. Compare CNC aluminum before committing."),
        ("Actuator/gearbox enclosure", "PA12-CF shell plus metal bearing seats", "Treat gear mesh, bearings, and threaded interfaces as metal/precision features."),
        ("Feet and impact pads", "TPU 95A", "Replaceable tread geometry; use PEBA only when rebound is intentionally desired."),
        ("Covers and guards", "ASA outdoors; ABS indoors", "Use SLA only when fine finish or visual verification outweighs impact and UV limits."),
    ]
    add_table(doc, ["Robot component", "Recommended material", "Release condition"], final_rows, [1.60, 1.68, 3.17], font_size=8.6, first_col_bold=True)

    add_label_paragraph(
        doc,
        "Bottom line: ",
        "For the next weight-reduction body and leg build, choose PA12-HP MJF as the baseline and PA12-CF only as a targeted stiffness experiment. Use PA11-HP where cracking or impact is the failure mode. Keep the feet flexible and keep high-concentration bearing, shaft, and actuator loads in metal interfaces.",
        fill=PALE_ORANGE,
    )

    # Sources
    doc.add_heading("Sources and Data Boundaries", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "All service details below were checked on 26 August 2026. JLC3DP pages can change without notice. "
        "Mechanical values are reproduced or summarized from supplier pages and should be verified against the current quote-linked technical data sheet."
    )
    sources = [
        ("JLC3DP 3D Printing Capabilities", "https://jlc3dp.com/capabilities/3d-printing-service", "Processes, material families, tolerances, finishing, and general design rules."),
        ("JLC3DP 3D Printing Material Catalog", "https://jlc3dp.com/help/catalog/123", "Current public catalog names for resin, nylon, plastic, and metal."),
        ("JLC3DP Design Guideline", "https://jlc3dp.com/help/article/3d-printing-design-guideline", "Build sizes, wall-thickness guidance, holes, and process tolerances."),
        ("Functional 3D Printing Guide", "https://jlc3dp.com/blog/functional-3d-printing", "Process behavior, functional-part boundaries, and validation cautions."),
        ("PA12-HP Nylon", "https://jlc3dp.com/help/article/pa12-hp-nylon", "MJF PA12 price floor, dimensions, HDT, tensile, flexural, impact, and elongation."),
        ("PA11-HP Nylon", "https://jlc3dp.com/help/article/pa11-hp-nylon", "MJF PA11 toughness, equal XY/Z data, elongation, and dual-load HDT."),
        ("3301PA Nylon", "https://jlc3dp.com/help/article/3301pa-nylon", "SLS 3301PA mechanical properties and limits."),
        ("3201PA-F Nylon", "https://jlc3dp.com/help/article/3201pa-f-nylon", "SLS 3201PA-F mechanical properties and print rules."),
        ("Precimid 1172 Pro Nylon", "https://jlc3dp.com/help/article/precimid-1172-pro", "SLS 1172 Pro mechanical properties and build envelope."),
        ("3401GB Nylon", "https://jlc3dp.com/help/article/3401gb-nylon", "Glass-bead nylon stiffness, strength, impact, elongation, and HDT."),
        ("PA12-CF Plastic", "https://jlc3dp.com/help/article/pa12-cf-plastic", "FDM PA12-CF strength, stiffness, elongation, HDT, and service rules."),
        ("PEEK Plastic", "https://jlc3dp.com/help/article/peek-plastic", "FDM PEEK XY/Z mechanics, thermal properties, and manufacturing limitations."),
        ("ABS Plastic", "https://jlc3dp.com/help/article/abs-plastic", "FDM ABS mechanics, HDT, cost floor, build size, and default infill."),
        ("ASA Plastic", "https://jlc3dp.com/help/article/asa-plastic", "FDM ASA UV use, mechanics, HDT, and service rules."),
        ("TPU", "https://jlc3dp.com/help/article/tpu", "FDM TPU hardness, applications, infill, and dimensions."),
        ("PEBA Plastic", "https://jlc3dp.com/help/article/peba-plastic", "FDM PEBA elasticity, tensile strength, modulus, density, and low-temperature behavior."),
        ("316L Stainless Steel", "https://jlc3dp.com/help/article/316l-stainless-steel", "SLM 316L strength, yield, elongation, modulus, roughness, and service rules."),
        ("Titanium TC4", "https://jlc3dp.com/help/article/titanium-tc4", "SLM TC4 strength, yield, elongation, and manufacturing limitations."),
        ("BJ-316L Stainless Steel", "https://jlc3dp.com/help/article/bj-316l-stainless-steel", "Binder-jet 316L strength, yield, elongation, tolerance, and build envelope."),
        ("JLC3DP Material Pricing Update", "https://jlc3dp.com/news/drastic-material-price-cuts-and-adjusted-post-processing-service-fees", "Published 2025 PA12 and SLM-316L price-per-gram changes; final quotes still depend on geometry and finishing."),
    ]
    for idx, (title, url, note) in enumerate(sources, 1):
        add_source(doc, idx, title, url, note)

    doc.add_heading("Important limitations", level=2)
    add_bullet(doc, "'From' prices are marketing entry points, not normalized $/cm³ or final part prices. Shipping, minimum charge, geometry, packing, surface finish, color, batch volume, and special handling can dominate.")
    add_bullet(doc, "Published strength may describe supplier test coupons rather than your exact infill, wall count, orientation, conditioning, or post-processing.")
    add_bullet(doc, "HDT is not continuous-use temperature and is strongly load-dependent.")
    add_bullet(doc, "Fatigue, creep, joint backlash, bearing-seat wear, and impact performance must be measured on representative robot parts.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
